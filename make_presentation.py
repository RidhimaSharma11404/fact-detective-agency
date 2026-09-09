import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG_COLOR = RGBColor(15, 23, 42)        # Slate 900
CARD_BG = RGBColor(30, 41, 59)         # Slate 800
CARD_BORDER = RGBColor(51, 65, 85)     # Slate 700
ACCENT_AMBER = RGBColor(245, 158, 11)  # Amber 500
ACCENT_GOLD = RGBColor(251, 191, 36)   # Amber 400
ACCENT_SKY = RGBColor(56, 189, 248)    # Sky 400
ACCENT_EMERALD = RGBColor(52, 211, 153)# Emerald 400
ACCENT_ROSE = RGBColor(244, 63, 94)    # Rose 500
ACCENT_INDIGO = RGBColor(129, 140, 248)# Indigo 400
TEXT_WHITE = RGBColor(248, 250, 252)   # Slate 50
TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, tag_text, title_text, category_color=ACCENT_AMBER):
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_right = tf_tag.margin_top = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = category_color

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

def create_card(slide, left, top, width, height, title, items, badge_text='', badge_color=ACCENT_AMBER, bg_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)
    
    p0 = tf.paragraphs[0]
    p0.text = f'{badge_text}  {title}' if badge_text else title
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = badge_color
    p0.space_after = Pt(10)
    
    for it in items:
        p = tf.add_paragraph()
        p.text = f'• {it}'
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(6)
    return shape

# SLIDE 1: Title
slide_layout = prs.slide_layouts[6]
s1 = prs.slides.add_slide(slide_layout)
set_slide_background(s1)

main_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5))
main_card.fill.solid()
main_card.fill.fore_color.rgb = CARD_BG
main_card.line.color.rgb = ACCENT_AMBER
main_card.line.width = Pt(2.0)

tf1 = main_card.text_frame
tf1.word_wrap = True
tf1.margin_left = Inches(0.6)
tf1.margin_right = Inches(0.6)
tf1.margin_top = Inches(0.6)

p_tag = tf1.paragraphs[0]
p_tag.text = 'MULTI-AGENT KNOWLEDGE GRAPH & ADJUDICATION SYSTEM'
p_tag.font.size = Pt(13)
p_tag.font.bold = True
p_tag.font.color.rgb = ACCENT_GOLD
p_tag.space_after = Pt(12)

p_h1 = tf1.add_paragraph()
p_h1.text = 'The Fact Detective Agency 🕵️‍♂️⚖️'
p_h1.font.size = Pt(36)
p_h1.font.bold = True
p_h1.font.color.rgb = TEXT_WHITE
p_h1.space_after = Pt(14)

p_sub = tf1.add_paragraph()
p_sub.text = 'Automated Fact Extraction, Verifiable Evidence Provenance & Cross-Document Adjudication Layer for Complex Enterprise Disclosures'
p_sub.font.size = Pt(16)
p_sub.font.color.rgb = TEXT_MUTED
p_sub.space_after = Pt(28)

p_links = tf1.add_paragraph()
p_links.text = '🚀 Live Web App: https://fact-detective-agency.loca.lt  |  ⚡ Vercel Production  |  🐙 GitHub: RidhimaSharma11404/fact-detective-agency'
p_links.font.size = Pt(12)
p_links.font.bold = True
p_links.font.color.rgb = ACCENT_SKY

# SLIDE 2: Problem
s2 = prs.slides.add_slide(slide_layout)
set_slide_background(s2)
add_header(s2, '01 / Problem Statement', 'Why Standard LLMs & Basic RAG Fail on Corporate Disclosures')

create_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), 'Hallucination & Provenance', [
    'Numbers and metrics extracted without exact character-level source citations.',
    'No verifiable bounding blocks or page-level evidence audit trail.',
    'Unchecked confidence scores masquerading as ground-truth facts.'
], '❌', ACCENT_ROSE)

create_card(s2, Inches(4.86), Inches(1.8), Inches(3.6), Inches(4.8), 'Temporal Blindness', [
    'Fails to distinguish across chronological reporting horizons.',
    'Pre-IPO baseline revenue confused with post-IPO consolidated scale.',
    'Historical facts treated as current state without supersession tracking.'
], '❌', ACCENT_AMBER)

create_card(s2, Inches(8.93), Inches(1.8), Inches(3.6), Inches(4.8), 'Scope & Line-Item Ambiguity', [
    'Directly equates "customer contract revenue" with "total operational income".',
    'Generates high false-alarm contradiction rates on legitimate accounting distinctions.',
    'Lacks structured epistemic modality calibration.'
], '❌', ACCENT_SKY)

# SLIDE 3: Architecture
s3 = prs.slides.add_slide(slide_layout)
set_slide_background(s3)
add_header(s3, '02 / System Architecture', 'Coordinated Multi-Agent Intelligence Pipeline')

create_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 'Agent 1: The Detective', [
    'Layout-Aware PDF Block Parser (PyMuPDF + PyPDF fallback).',
    'Extracts atomic <Subject, Predicate, Object> factual propositions.',
    'Dynamic Qualifier Inference (fiscal year, accounting basis, currency units).',
    'Epistemic Credibility Calibration (0.0 to 1.0) using certainty markers.',
    'Exact text span fuzzy matching with character-level offsets.'
], '🕵️‍♂️', ACCENT_GOLD)

create_card(s3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), 'Agent 2: The Courtroom Panel', [
    'Dense 256d Vector Index for O(K) sub-linear candidate pair retrieval.',
    'The Skeptic Judge: Adversarial prosecutor targeting numerical divergence.',
    'The Reconciler Judge: Contextual judge resolving scope & temporal horizons.',
    'Chief Magistrate Arbiter: Authoritative tiebreaker on epistemic deadlocks.',
    'Deterministic SHA-256 Fact Lifecycle & Supersession Chains.'
], '⚖️', ACCENT_SKY)

# SLIDE 4: Agent 1 Deep Dive
s4 = prs.slides.add_slide(slide_layout)
set_slide_background(s4)
add_header(s4, '03 / Extraction & Grounding', 'Agent 1: Atomic Claim Extraction & Epistemic Calibration')

create_card(s4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), 'Structured Claim Schema', [
    'Subject: Entity / Corporate body',
    'Predicate: Action / Reported metric',
    'Object: Numerical value or status',
    'Qualifiers: Dynamic key-value context',
    'Modality: Audited statutory vs forward guidance',
    'Deterministic Hash: SHA-256 fingerprint'
], '📌', ACCENT_AMBER)

create_card(s4, Inches(4.86), Inches(1.8), Inches(3.6), Inches(4.8), 'Epistemic Calibration (0-1)', [
    '0.90 – 1.00: Audited statutory balance sheets & official tables.',
    '0.70 – 0.89: Narrative management commentary & operational notes.',
    '0.50 – 0.69: Hedged or estimated estimates ("approximately").',
    '0.20 – 0.49: Forward guidance & management projections.'
], '🎯', ACCENT_EMERALD)

create_card(s4, Inches(8.93), Inches(1.8), Inches(3.6), Inches(4.8), 'Verbatim Evidence Grounding', [
    'Binds exact PDF document name.',
    'Records 1-indexed page number.',
    'Extracts exact character start & end offsets.',
    'Sliding-window token fuzzy matching with SequenceMatcher fallback.'
], '🔍', ACCENT_SKY)

# SLIDE 5: Agent 2 Deep Dive
s5 = prs.slides.add_slide(slide_layout)
set_slide_background(s5)
add_header(s5, '04 / Adversarial Adjudication', 'Agent 2: Dual-Judge NLI Debate & Chief Magistrate')

create_card(s5, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), 'The Skeptic Judge 🗡️', [
    'Adversarial prosecutor persona.',
    'Hunts for mathematical contradictions.',
    'Flags incompatible metric claims.',
    'Prevents single-prompt confirmation bias.'
], 'PROSECUTOR', ACCENT_ROSE)

create_card(s5, Inches(4.86), Inches(1.8), Inches(3.6), Inches(4.8), 'The Reconciler Judge ⚖️', [
    'Investigative contextual arbiter.',
    'Checks Temporal Intervals (FY21 vs FY24).',
    'Checks Line-Item Scope (contracts vs total).',
    'Checks Accounting Standards & Units.'
], 'RECONCILER', ACCENT_SKY)

create_card(s5, Inches(8.93), Inches(1.8), Inches(3.6), Inches(4.8), 'Chief Magistrate ⚡', [
    'Resolves HUNG_JURY deadlocks.',
    'Manages Fact Supersession lifecycles.',
    'Maintains unbroken audit provenance.',
    'Human-in-the-loop review overrides.'
], 'ARBITER', ACCENT_GOLD)

# SLIDE 6: Benchmark Exhibits
s6 = prs.slides.add_slide(slide_layout)
set_slide_background(s6)
add_header(s6, '05 / Curated Benchmarks', '5 Real-World Enterprise Adjudication Exhibits')

create_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.25), 'Case 1: Corroborated Statutory Revenue', [
    'Delhivery FY24 revenue confirmed across Annual Report & Presentation (₹81,420M).',
    'Verdict: CORROBORATED | Credibility: 0.95 | Verbatim evidence matched.'
], '🟢', ACCENT_EMERALD)

create_card(s6, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.25), 'Case 2: Direct Contradiction Flagged', [
    'Identical scope & temporal parameters with conflicting financial metrics.',
    'Verdict: CONTRADICTED | The Skeptic catches irreconcilable divergence.'
], '🔴', ACCENT_ROSE)

create_card(s6, Inches(0.8), Inches(4.35), Inches(3.6), Inches(2.35), 'Case 3: Scope Reconciliation', [
    'Contracts (₹68,810M) vs Total Income (₹72,250M).',
    'Verdict: RECONCILED (Zero False Alarms).'
], '🟣', ACCENT_INDIGO)

create_card(s6, Inches(4.86), Inches(4.35), Inches(3.6), Inches(2.35), 'Case 4: Supersession Lifecycle', [
    'Director status: FY22 Active -> FY24 Retired.',
    'Verdict: SUPERSEDED_BY with audit trail.'
], '🔵', ACCENT_SKY)

create_card(s6, Inches(8.93), Inches(4.35), Inches(3.6), Inches(2.35), 'Case 5: Hung Jury Arbiter', [
    'Hedged guidance vs audited loss deadlock.',
    'Resolved via Chief Magistrate tiebreaker.'
], '🟡', ACCENT_GOLD)

# SLIDE 7: Reliability & REST API
s7 = prs.slides.add_slide(slide_layout)
set_slide_background(s7)
add_header(s7, '06 / Production Engineering', 'Dual-Engine Reliability, Testing & REST API Ecosystem')

create_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 'Enterprise Reliability & Tests', [
    'Frontier LLMs (Gemini 2.0 / GPT-4o) + Zero-Dependency Heuristic Fallback.',
    'Dynamic Model Discovery auto-negotiates active API tiers without 404 errors.',
    '256-dimensional semantic vector index prevents O(N^2) quadratic explosions.',
    '8/8 Automated Pytest Unit Tests passing in 0.71s.',
    'Serverless Vercel & Docker Containerized deployment.'
], '⚙️', ACCENT_EMERALD)

create_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), 'Complete REST API & Swagger Docs', [
    'POST /api/documents/upload: Multi-agent PDF ingestion & extraction.',
    'GET /api/facts & GET /api/facts/{id}: Fact Dossier & witness verdicts.',
    'POST /api/verdicts/{id}/tiebreak: Chief Magistrate arbitration.',
    'POST /api/verdicts/{id}/review: Human-in-the-loop review override.',
    'Interactive Swagger Docs live at /docs and /redoc.'
], '🌐', ACCENT_SKY)

# SLIDE 8: Summary
s8 = prs.slides.add_slide(slide_layout)
set_slide_background(s8)
add_header(s8, '07 / Summary & Access', 'Key Takeaways & Live Demonstration Links')

summary_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
summary_card.fill.solid()
summary_card.fill.fore_color.rgb = CARD_BG
summary_card.line.color.rgb = ACCENT_AMBER
summary_card.line.width = Pt(1.5)

tf8 = summary_card.text_frame
tf8.word_wrap = True
tf8.margin_left = Inches(0.5)
tf8.margin_right = Inches(0.5)
tf8.margin_top = Inches(0.4)

p_sum_h = tf8.paragraphs[0]
p_sum_h.text = 'The Fact Detective Agency — Production Summary'
p_sum_h.font.size = Pt(20)
p_sum_h.font.bold = True
p_sum_h.font.color.rgb = ACCENT_GOLD
p_sum_h.space_after = Pt(14)

points = [
    'Transforms unstructured corporate PDFs into verifiable, audited knowledge graphs.',
    'Dual-judge adversarial adjudication eliminates single-prompt hallucinations and false alarms.',
    'Verifiable provenance: Every single claim is anchored to exact character-level PDF citations.',
    'Zero-failure resilience: Seamless transition between frontier LLMs and offline heuristic engines.'
]
for pt in points:
    p = tf8.add_paragraph()
    p.text = f'✔  {pt}'
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_WHITE
    p.space_after = Pt(8)

p_foot = tf8.add_paragraph()
p_foot.text = '\n🌐 Live App: https://fact-detective-agency.loca.lt (Password: 182.72.39.9)\n🐙 GitHub: https://github.com/RidhimaSharma11404/fact-detective-agency'
p_foot.font.size = Pt(13)
p_foot.font.bold = True
p_foot.font.color.rgb = ACCENT_SKY

out_desktop = Path(r'C:\Users\Lenovo\OneDrive\Desktop\Fact_Detective_Agency_Presentation.pptx')
out_project = Path(r'C:\Users\Lenovo\.gemini\antigravity\scratch\fact-detective-agency\Fact_Detective_Agency_Presentation.pptx')

prs.save(str(out_desktop))
prs.save(str(out_project))

print('SUCCESS:', out_desktop.exists(), out_project.exists())

